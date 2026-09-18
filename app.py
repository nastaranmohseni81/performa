import os
import io
import json
import sqlite3
import jdatetime
from num2fawords import words
import openpyxl
from openpyxl.styles import Font
from flask import Flask, request, render_template_string, send_file, jsonify

app = Flask(__name__)
app.secret_key = "quotation_secret_key"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
DB_NAME = os.path.join(BASE_DIR, "quotation_db.db")

def init_db():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS clients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE,
            debt INTEGER DEFAULT 0
        )
    ''')
    cursor.execute('''
        CREATE TABLE IF NOT EXISTS settings (
            key TEXT PRIMARY KEY,
            value TEXT
        )
    ''')
    cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('last_invoice_num', '')")
    cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('last_template_path', '')")
    cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('last_parts_path', '')")
    conn.commit()
    conn.close()

init_db()

def get_clients_dict():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("SELECT name, debt FROM clients ORDER BY name ASC")
    rows = cursor.fetchall()
    conn.close()
    return {r[0]: r[1] for r in rows}

def get_next_invoice_num():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("SELECT value FROM settings WHERE key='last_invoice_num'")
    row = cursor.fetchone()
    conn.close()
    if row and row[0]:
        try:
            return str(int(row[0]) + 1)
        except ValueError:
            return row[0]
    return ""

def safe_write(sheet, cell_coordinate, value, number_format=None, is_bold=None, font_size=None, font_name="B Nazanin"):
    try:
        cell = sheet[cell_coordinate]
        target_cell = cell
        merged_cells = []

        if type(cell).__name__ == 'MergedCell':
            for merged_range in sheet.merged_cells.ranges:
                if cell_coordinate in merged_range:
                    target_cell = sheet.cell(row=merged_range.min_row, column=merged_range.min_col)
                    for r in range(merged_range.min_row, merged_range.max_row + 1):
                        for c in range(merged_range.min_col, merged_range.max_col + 1):
                            merged_cells.append(sheet.cell(row=r, column=c))
                    break
        else:
            merged_cells.append(target_cell)

        target_cell.value = value
        current_font = target_cell.font
        size_to_use = font_size if font_size is not None else (current_font.size if (current_font and current_font.size) else 11)
        bold_to_use = is_bold if is_bold is not None else (current_font.bold if current_font else False)
        color_to_use = current_font.color if current_font else None

        new_font = Font(name=font_name, size=size_to_use, bold=bold_to_use, color=color_to_use)

        for c_cell in merged_cells:
            c_cell.font = new_font
            if number_format:
                c_cell.number_format = number_format
    except Exception:
        pass

def get_real_cell_value(sheet, row, col):
    try:
        cell = sheet.cell(row=row, column=col)
        if type(cell).__name__ == 'MergedCell':
            for merged_range in sheet.merged_cells.ranges:
                if cell.coordinate in merged_range:
                    return sheet.cell(row=merged_range.min_row, column=merged_range.min_col).value
        return cell.value
    except Exception:
        return None

HTML_TEMPLATE = """
<!DOCTYPE html>
<html lang="fa" dir="rtl">
<head>
    <meta charset="UTF-8">
    <title>سیستم صدور پیش‌فاکتور رسمی شرکت</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.rtl.min.css" rel="stylesheet">
    <link href="https://cdn.jsdelivr.net/gh/rastikerdar/vazirmatn@v33.003/Vazirmatn-font-face.css" rel="stylesheet" type="text/css" />
    <style>
        body { font-family: 'Vazirmatn', sans-serif; background-color: #f4f6f9; color: #2f3640; padding: 20px; }
        .main-card { background: #ffffff; border-radius: 10px; box-shadow: 0 4px 15px rgba(0,0,0,0.05); padding: 25px; }
        .section-box { background-color: #f8f9fa; border: 1px solid #dcdde1; border-radius: 8px; padding: 15px; margin-bottom: 15px; }
        .form-label { font-weight: bold; color: #2c3e50; font-size: 0.9rem; }
        .btn-custom { font-weight: bold; border-radius: 6px; }
        .table-header { background-color: #2c3e50; color: white; }
    </style>
</head>
<body>
<div class="container-fluid max-w-7xl main-card">
    <h4 class="text-center mb-4 fw-bold text-dark">سیستم صدور پیش‌فاکتور رسمی شرکت</h4>
    
    <form id="invoiceForm" method="POST" action="/export" enctype="multipart/form-data">
        <div class="section-box">
            <div class="row g-2 align-items-center">
                <div class="col-md-1">
                    <label class="form-label">شماره:</label>
                    <input type="text" name="invoice_num" id="invoice_num" class="form-control text-center" value="{{ invoice_num }}" required>
                </div>
                <div class="col-md-2">
                    <label class="form-label">تاریخ:</label>
                    <input type="text" name="invoice_date" class="form-control text-center" value="{{ today_date }}" required>
                </div>
                <div class="col-md-3">
                    <label class="form-label">مرکز درمانی:</label>
                    <div class="input-group">
                        <input list="clientsList" name="client_name" id="clientInput" class="form-control" placeholder="انتخاب یا تایپ..." onchange="clientSelected()" required>
                        <datalist id="clientsList">
                            {% for name in clients.keys() %}
                            <option value="{{ name }}">
                            {% endfor %}
                        </datalist>
                        <button type="button" class="btn btn-success btn-custom" data-bs-toggle="modal" data-bs-target="#addClientModal">+ مرکز جدید</button>
                    </div>
                </div>
                <div class="col-md-2">
                    <label class="form-label">موضوع:</label>
                    <select name="subject" id="subjectSelect" class="form-select" onchange="subjectChanged()" required>
                        <option value=""></option>
                        <option value="تعمیر کمپرسور">تعمیر کمپرسور</option>
                        <option value="تعمیر ونتیلاتور">تعمیر ونتیلاتور</option>
                        <option value="تعمیر ونتیلاتور پرتابل">تعمیر ونتیلاتور پرتابل</option>
                        <option value="سرویس دوره‌ای">سرویس دوره‌ای</option>
                        <option value="فروش قطعه">فروش قطعه</option>
                    </select>
                </div>
                <div class="col-md-2">
                    <label class="form-label">سریال:</label>
                    <input type="text" name="serial_num" id="serialInput" class="form-control" placeholder="سریال دستگاه">
                </div>
                <div class="col-md-2">
                    <label class="form-label">بدهی (ریال):</label>
                    <input type="number" name="debt" id="debtInput" class="form-control" value="0">
                </div>
            </div>
        </div>

        <div class="section-box">
            <div class="row g-2 align-items-center">
                <div class="col-md-4">
                    <button type="button" class="btn btn-primary btn-custom w-100" style="background-color: #9b59b6; border:none;" onclick="document.getElementById('partsExcelInput').click()">بارگذاری بانک قطعات (Excel)</button>
                    <input type="file" id="partsExcelInput" class="d-none" accept=".xlsx" onchange="loadPartsBank(this)">
                </div>
                <div class="col-md-8">
                    <span id="lblImportStatus" class="text-muted italic">لطفاً فایل اکسل قطعات را وارد کنید.</span>
                </div>
            </div>
        </div>

        <div class="section-box">
            <div class="row g-2 align-items-center">
                <div class="col-md-3">
                    <label class="form-label">نوع دستگاه:</label>
                    <select id="cmbDeviceType" class="form-select" onchange="deviceTypeChanged()">
                        <option value="">-- انتخاب دستگاه --</option>
                    </select>
                </div>
                <div class="col-md-4">
                    <label class="form-label">نام قطعه:</label>
                    <input list="partsList" id="cmbPart" class="form-control" placeholder="تایپ یا انتخاب قطعه..." oninput="partSelected()">
                    <datalist id="partsList"></datalist>
                </div>
                <div class="col-md-1">
                    <label class="form-label">تعداد:</label>
                    <input type="number" id="txtQuantity" class="form-control text-center" value="1">
                </div>
                <div class="col-md-2">
                    <label class="form-label">قیمت واحد:</label>
                    <input type="number" id="txtPartPrice" class="form-control" placeholder="ریال">
                </div>
                <div class="col-md-2 align-self-end">
                    <button type="button" class="btn btn-warning btn-custom w-100 text-white" style="background-color: #e67e22;" onclick="addItemToTable()">افزودن به جدول</button>
                </div>
            </div>
        </div>

        <table class="table table-bordered text-center align-middle" id="partsTable">
            <thead class="table-header">
                <tr>
                    <th style="width: 50px;">ردیف</th>
                    <th>نام قطعه / شرح دپارتمان</th>
                    <th style="width: 90px;">تعداد</th>
                    <th style="width: 160px;">قیمت واحد (ریال)</th>
                    <th style="width: 160px;">مبلغ کل (ریال)</th>
                    <th style="width: 80px;">عملیات</th>
                </tr>
            </thead>
            <tbody>
            </tbody>
        </table>

        <div class="section-box">
            <div class="row g-2 align-items-center">
                <div class="col-md-2">
                    <label class="form-label">مبلغ اجرت:</label>
                    <input type="number" name="wage" id="txtWage" class="form-control" value="0" oninput="calculateTotals()">
                </div>
                <div class="col-md-1">
                    <label class="form-label">تعداد اجرت:</label>
                    <input type="number" name="wage_qty" id="txtWageQty" class="form-control text-center" value="1" oninput="calculateTotals()">
                </div>
                <div class="col-md-2">
                    <label class="form-label">ایاب و ذهاب:</label>
                    <input type="number" name="transport" id="txtTransport" class="form-control" value="0" oninput="calculateTotals()">
                </div>
                <div class="col-md-7 d-flex justify-content-around align-items-center fw-bold fs-6">
                    <span id="lblSumParts">جمع قطعات: 0 ریال</span>
                    <span id="lblVat">ارزش افزوده (۱۰٪): 0 ریال</span>
                    <span id="lblTotal" class="text-success fs-5">قابل پرداخت: 0 ریال</span>
                </div>
            </div>
        </div>

        <div class="section-box">
            <label class="form-label">انتخاب فایل پیش‌فاکتور قالب رسمی شرکت (.xlsx):</label>
            <input type="file" name="template_file" class="form-control" accept=".xlsx" required>
        </div>

        <input type="hidden" name="items_json" id="itemsJson">

        <div class="row g-2 mt-3">
            <div class="col-md-9">
                <button type="submit" class="btn btn-success btn-custom w-100 py-3 fs-5" onclick="prepareSubmit()">صدور پیش‌فاکتور نهایی مطابق قالب رسمی (Excel)</button>
            </div>
            <div class="col-md-3">
                <button type="button" class="btn btn-warning btn-custom w-100 py-3 fs-5 text-white" onclick="resetForm()">ثبت و فاکتور جدید ➕</button>
            </div>
        </div>
    </form>
</div>

<div class="modal fade" id="addClientModal" tabindex="-1">
    <div class="modal-dialog">
        <div class="modal-content">
            <div class="modal-header">
                <h5 class="modal-title fw-bold">افزودن مرکز درمانی جدید</h5>
                <button type="button" class="btn-close" data-bs-dismiss="modal"></button>
            </div>
            <div class="modal-body">
                <div class="mb-3">
                    <label class="form-label">نام مرکز درمانی جدید:</label>
                    <input type="text" id="newClientName" class="form-control" placeholder="مثال: بیمارستان 17 شهریور رشت">
                </div>
                <div class="mb-3">
                    <label class="form-label">بدهی پیشین (ریال):</label>
                    <input type="number" id="newClientDebt" class="form-control" value="0" placeholder="میزان بدهی پیشین">
                </div>
            </div>
            <div class="modal-footer">
                <button type="button" class="btn btn-success btn-custom w-100" onclick="saveClientAjax()">ذخیره مرکز درمانی</button>
            </div>
        </div>
    </div>
</div>

<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script>
<script src="https://cdn.jsdelivr.net/npm/xlsx@0.18.5/dist/xlsx.full.min.js"></script>
<script>
let clientsData = {{ clients | tojson }};
let partsData = [];
let items = [];

function clientSelected() {
    let clientName = document.getElementById('clientInput').value;
    if (clientsData[clientName] !== undefined) {
        document.getElementById('debtInput').value = clientsData[clientName];
    }
}

function saveClientAjax() {
    let name = document.getElementById('newClientName').value.trim();
    let debt = document.getElementById('newClientDebt').value.trim() || "0";

    if (!name) { alert("نام مرکز نمی‌تواند خالی باشد."); return; }

    fetch('/add_client', {
        method: 'POST',
        headers: { 'Content-Type': 'application/x-www-form-urlencoded' },
        body: `name=${encodeURIComponent(name)}&debt=${encodeURIComponent(debt)}`
    })
    .then(r => r.json())
    .then(data => {
        if (data.status === 'success') {
            clientsData[name] = parseInt(debt);
            let datalist = document.getElementById('clientsList');
            let opt = document.createElement('option');
            opt.value = name;
            datalist.appendChild(opt);
            
            document.getElementById('clientInput').value = name;
            document.getElementById('debtInput').value = debt;
            
            let modal = bootstrap.Modal.getInstance(document.getElementById('addClientModal'));
            modal.hide();
            document.getElementById('newClientName').value = "";
            document.getElementById('newClientDebt').value = "0";
        } else {
            alert(data.message);
        }
    });
}

function loadPartsBank(input) {
    if (!input.files.length) return;
    let file = input.files[0];
    let reader = new FileReader();

    reader.onload = function(e) {
        let data = new Uint8Array(e.target.result);
        let workbook = XLSX.read(data, {type: 'array'});
        let firstSheet = workbook.Sheets[workbook.SheetNames[0]];
        let jsonData = XLSX.utils.sheet_to_json(firstSheet, {header: 1});

        partsData = [];
        let devices = new Set();

        jsonData.forEach((row, index) => {
            if (index === 0 || !row[1]) return;
            let code = row[0];
            let partName = String(row[1]).trim();
            let price = parseInt(row[2]) || 0;
            let deviceType = row[3] ? String(row[3]).trim() : "نامشخص";

            partsData.push({ code: code, name: partName, price: price, device: deviceType });
            if (deviceType) devices.add(deviceType);
        });

        let deviceSelect = document.getElementById('cmbDeviceType');
        deviceSelect.innerHTML = '<option value="">-- انتخاب دستگاه --</option>';
        Array.from(devices).sort().forEach(dev => {
            let opt = document.createElement('option');
            opt.value = dev;
            opt.innerText = dev;
            deviceSelect.appendChild(opt);
        });

        document.getElementById('lblImportStatus').innerText = `بانک قطعات با موفقیت بروز شد (${partsData.length} کالا)`;
    };
    reader.readAsArrayBuffer(file);
}

function deviceTypeChanged() {
    let selectedDevice = document.getElementById('cmbDeviceType').value;
    let datalist = document.getElementById('partsList');
    datalist.innerHTML = "";

    let filtered = partsData.filter(p => p.device === selectedDevice);
    filtered.forEach(p => {
        let opt = document.createElement('option');
        opt.value = p.name;
        datalist.appendChild(opt);
    });
}

function partSelected() {
    let partName = document.getElementById('cmbPart').value.trim();
    let found = partsData.find(p => p.name === partName);
    if (found) {
        document.getElementById('txtPartPrice').value = found.price;
    }
}

function subjectChanged() {
    let subject = document.getElementById('subjectSelect').value;
    let serialInput = document.getElementById('serialInput');
    if (subject === "سرویس دوره‌ای") {
        serialInput.value = "";
        serialInput.disabled = true;
    } else {
        serialInput.disabled = false;
    }
}

function addItemToTable() {
    let partName = document.getElementById('cmbPart').value.trim();
    let qty = parseInt(document.getElementById('txtQuantity').value) || 1;
    let price = parseInt(document.getElementById('txtPartPrice').value) || 0;

    if (!partName) { alert("لطفاً نام قطعه را وارد کنید."); return; }

    items.push({ name: partName, qty: qty, price: price, total: qty * price });
    document.getElementById('cmbPart').value = "";
    document.getElementById('txtQuantity').value = "1";
    document.getElementById('txtPartPrice').value = "";
    renderTable();
}

function deleteRow(index) {
    items.splice(index, 1);
    renderTable();
}

function renderTable() {
    let tbody = document.querySelector('#partsTable tbody');
    tbody.innerHTML = "";
    items.forEach((item, index) => {
        let tr = `<tr>
            <td>${index + 1}</td>
            <td class="text-start">${item.name}</td>
            <td>${item.qty}</td>
            <td>${item.price.toLocaleString()}</td>
            <td>${item.total.toLocaleString()}</td>
            <td><button type="button" class="btn btn-outline-danger btn-sm" onclick="deleteRow(${index})">حذف</button></td>
        </tr>`;
        tbody.innerHTML += tr;
    });
    calculateTotals();
}

function calculateTotals() {
    let sumParts = items.reduce((acc, curr) => acc + curr.total, 0);
    let wageUnit = parseInt(document.getElementById('txtWage').value) || 0;
    let wageQty = parseInt(document.getElementById('txtWageQty').value) || 1;
    let wage = wageUnit * wageQty;
    let transport = parseInt(document.getElementById('txtTransport').value) || 0;

    let taxable = sumParts + wage;
    let vat = Math.round(taxable * 0.10);
    let payable = taxable + vat + transport;

    document.getElementById('lblSumParts').innerText = "جمع قطعات: " + sumParts.toLocaleString() + " ریال";
    document.getElementById('lblVat').innerText = "ارزش افزوده (۱۰٪): " + vat.toLocaleString() + " ریال";
    document.getElementById('lblTotal').innerText = "قابل پرداخت: " + payable.toLocaleString() + " ریال";
}

function resetForm() {
    items = [];
    renderTable();
    document.getElementById('txtWage').value = "0";
    document.getElementById('txtWageQty').value = "1";
    document.getElementById('txtTransport').value = "0";
    document.getElementById('serialInput').value = "";
}

function prepareSubmit() {
    document.getElementById('itemsJson').value = JSON.stringify(items);
}
</script>
</body>
</html>
"""

@app.route("/")
def index():
    today_date = jdatetime.date.today().strftime("%Y/%m/%d")
    invoice_num = get_next_invoice_num()
    clients = get_clients_dict()
    return render_template_string(HTML_TEMPLATE, today_date=today_date, invoice_num=invoice_num, clients=clients)

@app.route("/add_client", methods=["POST"])
def add_client():
    name = request.form.get("name", "").strip()
    debt = int(request.form.get("debt", 0) or 0)
    if not name:
        return jsonify({"status": "error", "message": "نام مرکز نمی‌تواند خالی باشد."})
        
    try:
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("INSERT INTO clients (name, debt) VALUES (?, ?)", (name, debt))
        conn.commit()
        conn.close()
        return jsonify({"status": "success"})
    except sqlite3.IntegrityError:
        return jsonify({"status": "error", "message": "این مرکز درمانی قبلاً ثبت شده است."})

@app.route("/export", methods=["POST"])
def export():
    invoice_num = request.form.get("invoice_num", "").strip()
    invoice_date = request.form.get("invoice_date", "").strip()
    client_name = request.form.get("client_name", "").strip()
    subject_selected = request.form.get("subject", "").strip()
    serial_num = request.form.get("serial_num", "").strip()
    debt_val = int(request.form.get("debt") or 0)
    
    wage_unit = int(request.form.get("wage") or 0)
    wage_qty = int(request.form.get("wage_qty") or 1)
    trans_val = int(request.form.get("transport") or 0)
    
    items_raw = request.form.get("items_json", "[]")
    items = json.loads(items_raw)
    
    template_file = request.files.get("template_file")
    if not template_file:
        return "لطفاً فایل قالب اکسل را آپلود کنید", 400

    if subject_selected == "سرویس دوره‌ای":
        full_subject = "سرویس دوره‌ای"
    else:
        full_subject = f"{subject_selected} به شماره سریال {serial_num}" if serial_num else subject_selected

    wb = openpyxl.load_workbook(template_file)
    sheet = wb.active

    safe_write(sheet, 'J2', f"شماره: {invoice_num}")
    safe_write(sheet, 'J3', f"تاریخ: {invoice_date}")
    safe_write(sheet, 'C10', full_subject, is_bold=True)
    safe_write(sheet, 'C11', client_name, is_bold=True)

    start_row = 13
    default_rows = 13

    all_entries = []
    for item in items:
        all_entries.append({
            "name": item["name"],
            "price": item["price"],
            "qty": item["qty"],
            "unit": "عدد",
            "total": item["total"]
        })

    if wage_unit > 0:
        all_entries.append({
            "name": "اجرت",
            "price": wage_unit,
            "qty": wage_qty,
            "unit": "عدد",
            "total": wage_unit * wage_qty
        })
    if trans_val > 0:
        all_entries.append({
            "name": "هزینه ایاب و ذهاب",
            "price": trans_val,
            "qty": 1,
            "unit": "عدد",
            "total": trans_val
        })

    active_rows_count = max(len(all_entries), default_rows)
    for i in range(active_rows_count):
        row_idx = start_row + i
        if i < len(all_entries):
            item = all_entries[i]
            safe_write(sheet, f'B{row_idx}', i + 1)
            safe_write(sheet, f'C{row_idx}', item["name"], is_bold=False)
            safe_write(sheet, f'G{row_idx}', item["price"], number_format='#,##0')
            safe_write(sheet, f'H{row_idx}', item["qty"])
            safe_write(sheet, f'I{row_idx}', item["unit"])
            safe_write(sheet, f'J{row_idx}', item["total"], number_format='#,##0')
        else:
            safe_write(sheet, f'B{row_idx}', "")
            safe_write(sheet, f'C{row_idx}', "")
            safe_write(sheet, f'G{row_idx}', "")
            safe_write(sheet, f'H{row_idx}', "")
            safe_write(sheet, f'I{row_idx}', "")
            safe_write(sheet, f'J{row_idx}', "")

    sum_parts = sum(item["total"] for item in items)
    wage = wage_unit * wage_qty
    taxable = sum_parts + wage
    vat = int(taxable * 0.10)
    payable = taxable + vat + trans_val

    sum_row = 26
    vat_row = 27
    payable_row = 28

    safe_write(sheet, f'J{sum_row}', sum_parts + wage, number_format='#,##0')
    safe_write(sheet, f'J{vat_row}', vat, number_format='#,##0')
    safe_write(sheet, f'J{payable_row}', payable, number_format='#,##0')

    payable_in_words = words(payable) + " ریال"
    safe_write(sheet, f'B{sum_row}', f"مبلغ پرداختی به حروف: {payable_in_words}")

    formatted_debt = f"{debt_val:,}"
    debt_written = False
    for row_num in range(29, sheet.max_row + 1):
        for col_num in range(1, sheet.max_column + 1):
            val_str = str(get_real_cell_value(sheet, row_num, col_num) or "")
            if val_str.strip() == "مبلغ":
                target_cell = sheet.cell(row=row_num, column=col_num + 1)
                safe_write(sheet, target_cell.coordinate, debt_val, number_format='#,##0', is_bold=True)
                debt_written = True
                break
            elif "استحضار" in val_str or "بدهی" in val_str:
                debt_text_full = (
                    f"به استحضار می‌رساند مجموع بدهی پیشین آن مرکز محترم بابت خدمات ارائه شده "
                    f"مبلغ {formatted_debt} ریال می‌باشد.\n"
                    f"خواهشمند است جهت تداوم خدمات مطلوب با قید فوریت نسبت به تسویه حساب بدهی فوق اقدام فرمایید."
                )
                cell_to_write = sheet.cell(row=row_num, column=col_num)
                safe_write(sheet, cell_to_write.coordinate, debt_text_full, is_bold=True)
                debt_written = True
                break
        if debt_written:
            break

    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("UPDATE settings SET value=? WHERE key='last_invoice_num'", (invoice_num,))
    cursor.execute("UPDATE clients SET debt=? WHERE name=?", (debt_val, client_name))
    conn.commit()
    conn.close()

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        download_name=f"{invoice_num} - {client_name}.xlsx",
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

if __name__ == "__main__":
    app.run()